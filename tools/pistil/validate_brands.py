#!/usr/bin/env python3
"""Verify every brand-filtered store_rank pull against a known-good total.

Pistil fails quietly in two ways that both produce plausible files:
  - an ignored ctrl_BRAND returns the FULL dataset (Trap 4)
  - too short a render exports the PREVIOUS query (Trap 7) — seen as 729 rows
    (NY + Michigan) and two brands sharing an identical total

Neither row count nor a flat volume ceiling catches both. The decisive test is the
brand's own statewide total from the brand-rank export: a correct store-rank cut
lands at ~92-93% of it (the remainder sits at doors below the rank cutoff).

  python validate_brands.py          # report
  python validate_brands.py --fix    # also move bad files aside for re-pull
"""
import os, sys, glob, re, collections
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
DL = os.path.join(os.path.expanduser("~"), "Downloads")
BRAND_RANK = os.path.join(DL, "PULL_NY_brand_1mo.xlsx")

LO, HI = 0.80, 1.10      # acceptable store-rank / brand-rank ratio
MAX_ROWS = 680           # NY has ~645 doors; above this, Michigan has leaked in
MAX_VOL = 20_000_000     # no NY brand is close; above this the filter was ignored


def totals(path):
    # read_only keeps the file handle open until closed — without this, --fix fails
    # with a Windows PermissionError and the bad files stay in place to be ingested.
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] is not None]
        return len(rows), sum(r[4] or 0 for r in rows)
    finally:
        wb.close()


def main():
    expected = {}
    if os.path.exists(BRAND_RANK):
        wb = openpyxl.load_workbook(BRAND_RANK, data_only=True, read_only=True)
        ws = wb[wb.sheetnames[0]]
        for r in ws.iter_rows(min_row=2, values_only=True):
            if r[0] is not None:
                expected[str(r[1]).strip().lower()] = r[3] or 0
    else:
        print(f"! {BRAND_RANK} missing — falling back to row/volume sanity only\n")

    real = {}
    lst = os.path.join(HERE, "brands.txt")
    if os.path.exists(lst):
        for line in open(lst, encoding="utf-8"):
            b = line.strip()
            if b:
                real[re.sub(r"_+", "_", re.sub(r"[^A-Za-z0-9]", "_", b)).strip("_")] = b

    files = sorted(glob.glob(os.path.join(DL, "PULL_BRAND_*.xlsx")))
    by_total = collections.defaultdict(list)
    bad, good = [], []

    for f in files:
        slug = os.path.basename(f)[len("PULL_BRAND_"):-len(".xlsx")]
        name = real.get(slug, slug.replace("_", " "))
        try:
            n, t = totals(f)
        except Exception as e:
            bad.append((name, f, f"unreadable: {e}"))
            continue
        by_total[round(t)].append(name)
        exp = expected.get(name.lower())
        why = None
        if n == 0:
            why = "0 rows"
        elif n > MAX_ROWS:
            why = f"{n} rows — Michigan contamination (>{MAX_ROWS})"
        elif t >= MAX_VOL:
            why = f"${t:,.0f} — brand filter ignored"
        elif exp:
            ratio = t / exp if exp else 0
            if not (LO <= ratio <= HI):
                why = f"${t:,.0f} vs brand-rank ${exp:,.0f} = {ratio:.0%}"
        (bad.append((name, f, why)) if why else good.append((name, n, t, exp)))

    # Two brands with byte-identical totals means one is a stale render of the other.
    dupes = {t: ns for t, ns in by_total.items() if len(ns) > 1}

    print(f"{len(good)} OK, {len(bad)} bad, {len(files)} files\n")
    for name, n, t, exp in good:
        r = f"{t/exp:.0%}" if exp else "  -"
        print(f"  OK   {name:<28} {n:>4} doors  ${t:>11,.0f}  {r:>4} of brand-rank")
    if dupes:
        print("\nIDENTICAL TOTALS (stale render — one is a copy of the other):")
        for t, ns in dupes.items():
            print(f"  ${t:>11,}  {', '.join(ns)}")
    if bad:
        print("\nBAD:")
        for name, f, why in bad:
            print(f"  BAD  {name:<28} {why}")
        if "--fix" in sys.argv:
            for name, f, why in bad:
                os.replace(f, f + ".rejected")
            print(f"\nMoved {len(bad)} file(s) aside — re-run pull_brands.sh to re-pull them.")
        else:
            print("\nRe-run with --fix to move these aside for re-pulling.")
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(main())
