#!/usr/bin/env python3
"""What each brand actually sells — from the Pistil product-rank export.

The bot needs to answer "I'm selling Green Revolution" without being told what
Green Revolution makes. The product-rank export already knows: every SKU carries
its brand, category, units, dollars and average price, because product names are
formatted `Brand - Strain - Size - Form`.

That makes the product offering measured rather than guessed — the bot only needs
to reach for web search to enrich positioning (brand story, retail pricing, new
launches), not to discover what the brand sells.

Writes server/brand_profiles.json:
  {brand: {vol, categories: [{cat, vol, share, avg_price, units}], top_products: [...]}}

Usage:
  python tools/build_brand_profiles.py [product_rank.xlsx]
"""
import os, sys, json, glob, collections
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "server", "brand_profiles.json")
DL = os.path.join(os.path.expanduser("~"), "Downloads")

# A product-rank export covering the whole market has tens of thousands of rows.
# A short one is a partial Sigma render (Trap 3) and would silently under-report
# every brand's offering, so refuse it rather than publish a thin profile.
MIN_ROWS = 20_000


def main():
    pos = [a for a in sys.argv[1:] if not a.startswith("--")]
    if pos:
        path = os.path.expanduser(pos[0])
    else:
        cands = glob.glob(os.path.join(DL, "PULL_NY_prod_*.xlsx"))
        if not cands:
            raise SystemExit("No PULL_NY_prod_*.xlsx in ~/Downloads — pull product_rank first.")
        path = max(cands, key=os.path.getsize)   # biggest = most complete render

    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        it = ws.iter_rows(values_only=True)
        next(it)
        rows = []
        for r in it:
            if r[0] is None:
                continue
            name = str(r[2] or "")
            brand = name.split(" - ")[0].strip()
            if not brand:
                continue
            rows.append((brand, str(r[1] or "?"), name,
                         r[3] if isinstance(r[3], (int, float)) else 0,
                         r[4] if isinstance(r[4], (int, float)) else 0,
                         r[8] if isinstance(r[8], (int, float)) else 0))
    finally:
        wb.close()

    if len(rows) < MIN_ROWS:
        raise SystemExit(
            f"{os.path.basename(path)} has only {len(rows):,} rows (expected >{MIN_ROWS:,}) — "
            "partial render, refusing to build thin brand profiles.")

    print(f"{os.path.basename(path)}: {len(rows):,} products")

    by_brand = collections.defaultdict(lambda: {"vol": 0.0, "cats": collections.defaultdict(
        lambda: {"vol": 0.0, "units": 0.0, "prices": []}), "prods": []})
    for brand, cat, name, units, vol, price in rows:
        b = by_brand[brand]
        b["vol"] += vol
        c = b["cats"][cat]
        c["vol"] += vol
        c["units"] += units
        # Avg Menu Price is per-SKU; a handful of rows carry junk values, so weight
        # by volume and drop absurd ones rather than letting one row skew the mean.
        if 0 < price < 1000:
            c["prices"].append((price, vol))
        b["prods"].append((name, vol))

    profiles = {}
    for brand, b in by_brand.items():
        if b["vol"] <= 0:
            continue
        cats = []
        for cat, c in b["cats"].items():
            if c["vol"] <= 0:
                continue
            wsum = sum(v for _, v in c["prices"]) or 1
            avg = sum(p * v for p, v in c["prices"]) / wsum if c["prices"] else 0
            cats.append({"cat": cat, "vol": round(c["vol"]),
                         "share": round(c["vol"] / b["vol"] * 100),
                         "units": round(c["units"]),
                         "avg_price": round(avg, 2)})
        cats.sort(key=lambda x: -x["vol"])
        top = sorted(b["prods"], key=lambda x: -x[1])[:8]
        profiles[brand] = {
            "vol": round(b["vol"]),
            "categories": cats,
            "top_products": [{"name": n, "vol": round(v)} for n, v in top],
        }

    json.dump(profiles, open(OUT, "w"), indent=1)
    print(f"wrote {len(profiles)} brand profiles -> {OUT}")
    for name in ("Dragonfly", "Jerome Baker", "Green Revolution"):
        p = profiles.get(name)
        if not p:
            print(f"  ! {name}: not found in the product export")
            continue
        cats = ", ".join(f"{c['cat']} {c['share']}% (${c['avg_price']})" for c in p["categories"][:4])
        print(f"  {name:<18} ${p['vol']:>10,}  {cats}")


if __name__ == "__main__":
    main()
